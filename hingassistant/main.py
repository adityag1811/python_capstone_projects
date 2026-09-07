import openpyxl
from excel_reader import get_ingredients,get_cost
from calculate_quantity import calculate_quantity

wb = openpyxl.load_workbook('formula.xlsx',data_only=True)
another_batch = True

while another_batch:
    valid_product = False
    while not valid_product:

        product_name = input("Enter product name: ")
        if product_name in wb.sheetnames:
            valid_product = True
        else:
            print("Please enter a valid product name")

    valid = False
    while not valid:
        try:
            user_required_qty = int(input("Enter required quantity: "))
            valid = True
        except:
            print("Please enter a valid number")

    product_sheet = wb[product_name]
    last_row = product_sheet.max_row
    batch_qty = product_sheet.cell(row=1, column=13).value

    required_qty = user_required_qty
    scaling_factor = required_qty/batch_qty

    ingredients = get_ingredients(product_sheet,last_row)
    product_costing = get_cost(product_sheet,last_row)
    step_1, step_2, step_3 = calculate_quantity(ingredients, scaling_factor)

    print(f"\nProduct : {product_name}")
    print(f"Product Costing : Rs.{round(product_costing, 2)}/-")
    print(f"Required Quantity : {required_qty}")
    print("\nStep 1")

    for key,value in step_1.items():
        print(f"{key}: {value} Kg")

    if step_2:
        print("\nStep 2")
        for key, value in step_2.items():
            print(f"{key}: {value} Kg")

    if step_3:
        print("\nStep 3")
        for key,value in step_3.items():
            print(f"{key}: {value} Kg")

    while True:
        another_batch_input = input("Would you like to calculate another batch (y/n)?: ")
        if another_batch_input!="y" and another_batch_input!="n":
            print("Please enter either 'y' or 'n'")
        elif another_batch_input== "y":
            another_batch = True
            break
        elif another_batch_input == "n":
            another_batch = False
            print("Thanks for using our program")
            break
