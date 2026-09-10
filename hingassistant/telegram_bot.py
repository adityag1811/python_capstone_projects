import os
import openpyxl

from dotenv import load_dotenv
from telegram import Update
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    ConversationHandler,
    ContextTypes,
    filters
)

from starlette.applications import Starlette
from starlette.requests import Request
from starlette.responses import PlainTextResponse
from starlette.routing import Route
from uvicorn import run


load_dotenv()

bot_token = os.getenv("TELEGRAM_BOT_TOKEN")

wb = openpyxl.load_workbook("formula.xlsx", data_only=True)

PRODUCT, QUANTITY = range(2)


application = (
    Application.builder()
    .token(bot_token)
    .updater(None)
    .build()
)


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):

    await update.message.reply_text(
        "Which product would you like to calculate?"
    )

    return PRODUCT


async def get_product(update: Update, context: ContextTypes.DEFAULT_TYPE):

    product_name = update.message.text

    if product_name not in wb.sheetnames:

        await update.message.reply_text(
            "Please enter a valid product name."
        )

        return PRODUCT

    context.user_data["product_name"] = product_name

    await update.message.reply_text(
        "How many kg would you like?"
    )

    return QUANTITY


async def get_quantity(update: Update, context: ContextTypes.DEFAULT_TYPE):

    try:
        required_qty = int(update.message.text)

    except ValueError:

        await update.message.reply_text(
            "Please enter a valid whole number."
        )

        return QUANTITY

    if required_qty <= 0:

        await update.message.reply_text(
            "Please enter a quantity greater than zero."
        )

        return QUANTITY

    product_name = context.user_data["product_name"]

    product_sheet = wb[product_name]

    last_row = product_sheet.max_row

    batch_qty = product_sheet.cell(
        row=1,
        column=13
    ).value

    scaling_factor = required_qty / batch_qty

    ingredients = get_ingredients(
        product_sheet,
        last_row
    )

    product_costing = get_cost(
        product_sheet,
        last_row
    )

    step_1, step_2, step_3 = calculate_quantity(
        ingredients,
        scaling_factor
    )

    response = f"""
Product: {product_name}
Required Quantity: {required_qty} kg

Step 1
"""

    for key, value in step_1.items():
        response += f"\n{key}: {value} Kg"

    if step_2:

        response += "\n\nStep 2"

        for key, value in step_2.items():
            response += f"\n{key}: {value} Kg"

    if step_3:

        response += "\n\nStep 3"

        for key, value in step_3.items():
            response += f"\n{key}: {value} Kg"

    response += f"\n\nNet Cost: ₹{product_costing} per kg"

    await update.message.reply_text(response)

    return ConversationHandler.END


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):

    await update.message.reply_text(
        "Calculation cancelled."
    )

    return ConversationHandler.END


conversation_handler = ConversationHandler(

    entry_points=[
        CommandHandler("start", start)
    ],

    states={

        PRODUCT: [
            MessageHandler(
                filters.TEXT & ~filters.COMMAND,
                get_product
            )
        ],

        QUANTITY: [
            MessageHandler(
                filters.TEXT & ~filters.COMMAND,
                get_quantity
            )
        ]

    },

    fallbacks=[
        CommandHandler("cancel", cancel)
    ]
)


application.add_handler(conversation_handler)


async def telegram_webhook(request: Request):

    data = await request.json()

    update = Update.de_json(
        data,
        application.bot
    )

    await application.update_queue.put(update)

    return PlainTextResponse("OK")


async def homepage(request: Request):

    return PlainTextResponse(
        "HingFormula Bot is running!"
    )


routes = [
    Route("/", homepage, methods=["GET"]),
    Route("/telegram", telegram_webhook, methods=["POST"])
]


app = Starlette(routes=routes)


if __name__ == "__main__":

    run(
        app,
        host="0.0.0.0",
        port=int(os.environ.get("PORT", 5000))
    )
